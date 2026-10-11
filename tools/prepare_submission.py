"""Add the real recording and Drive link; enforce the provided submission limits."""
from pathlib import Path
from urllib.parse import urlparse
import argparse
from copy import deepcopy
import hashlib
import json
import math
import shutil
import struct
import sys
import zipfile

NAME = 'DRL_AFrameworkForFairEvaluationOfVariance-AwareBanditAlgorithms_Group6'
BASE = Path(__file__).resolve().parents[1]
KIT = BASE / 'presentation'
SUB = BASE / 'submission'
LIMIT = 10_000_000

def atoms(f, start, end):
    pos = start
    while pos + 8 <= end:
        f.seek(pos)
        size, kind = struct.unpack('>I4s', f.read(8))
        header = 8
        if size == 1:
            size = struct.unpack('>Q', f.read(8))[0]
            header = 16
        elif size == 0:
            size = end - pos
        if size < header or pos + size > end:
            raise ValueError('Invalid MP4 atom size.')
        yield kind, pos + header, pos + size
        pos += size

def mp4_duration(path):
    with path.open('rb') as f:
        for kind, start, end in atoms(f, 0, path.stat().st_size):
            if kind != b'moov':
                continue
            for child, a, b in atoms(f, start, end):
                if child != b'mvhd':
                    continue
                f.seek(a)
                data = f.read(min(b-a, 40))
                if data[0] == 0:
                    _, _, scale, duration = struct.unpack_from('>IIII',data,4)
                elif data[0] == 1:
                    _, _, scale, duration = struct.unpack_from('>QQIQ',data,4)
                else:
                    raise ValueError('Unsupported MP4 movie-header version.')
                if not scale or not duration:
                    raise ValueError('Recording duration is missing or zero.')
                seconds = duration / scale
                if not math.isfinite(seconds):
                    raise ValueError('Recording duration is invalid.')
                return seconds
    raise ValueError('No movie duration found. Export the actual recording as MP4.')

def capture_open_source(source, snapshot):
    import win32com.client
    app = win32com.client.GetActiveObject('PowerPoint.Application')
    ending = '/presentation/' + source.name.lower()
    for p in app.Presentations:
        if p.FullName.replace('\\','/').lower().endswith(ending):
            p.SaveCopyAs(str(snapshot),24)
            return snapshot
    raise RuntimeError('Close the recording PPT and retry so it can be read.')

def set_drive_link(source, target, drive_link):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation(str(source))
    if len(prs.slides) != 16:
        raise ValueError('Expected the final 16-slide Group 6 deck.')
    hits = [s for s in prs.slides[-1].shapes if s.name == 'Group6_Recording_Link']
    if len(hits) != 1:
        raise ValueError('The recording-link box was not found on the last slide.')
    shape = hits[0]
    frame = shape.text_frame
    p = frame.paragraphs[0]
    style = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    props = p._p.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
    for node in props[1:]:
        p._p.remove(node)
    run = p.add_run()
    run.text = 'Video recording on Google Drive — click to watch'
    if style is not None:
        run._r.insert(0,style)
    run.font.size = Pt(17)
    for other in list(frame.paragraphs)[1:]:
        other._p.getparent().remove(other._p)
    shape.click_action.hyperlink.address = drive_link
    for r in p.runs:
        r.hyperlink.address = None
    prs.save(str(target))
    check = Presentation(str(target))
    last = next(s for s in check.slides[-1].shapes if s.name == 'Group6_Recording_Link')
    if last.click_action.hyperlink.address != drive_link:
        raise ValueError('The Drive hyperlink was not saved correctly.')

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--video',required=True,help='Path to your actual finished MP4.')
    parser.add_argument('--drive-link',required=True,help='Actual Google Drive viewing URL.')
    args = parser.parse_args()
    video = Path(args.video).expanduser().resolve()
    if not video.is_file() or video.suffix.lower() != '.mp4':
        raise ValueError('Supply the actual recorded video as an existing MP4 file.')
    if video.stat().st_size > LIMIT:
        raise ValueError('Video exceeds 10 MB. Re-export it within the assignment limit.')
    seconds = mp4_duration(video)
    if seconds > 420:
        raise ValueError(f'Video is {seconds:.2f} seconds; the maximum is 420 seconds.')
    link = args.drive_link.strip()
    parsed = urlparse(link)
    if parsed.scheme != 'https' or parsed.hostname != 'drive.google.com':
        raise ValueError('Provide an actual HTTPS Google Drive viewing link.')
    if any(x in link.upper() for x in ['REAL_ID','PASTE','PLACEHOLDER']):
        raise ValueError('Replace the example with the actual recording link.')
    if '/file/d/' not in parsed.path and 'id=' not in parsed.query:
        raise ValueError('Use the Drive link to the recorded video file.')
    SUB.mkdir(exist_ok=True)
    staged_ppt = SUB / (NAME + '.staged.pptx')
    staged_zip = SUB / (NAME + '.staged.zip')
    snapshot = SUB / 'open_recording_snapshot.pptx'
    staged = [staged_ppt, staged_zip, snapshot]
    try:
        source = KIT / (NAME + '.pptx')
        try:
            with source.open('rb') as f:
                f.read(1)
        except PermissionError:
            source = capture_open_source(source,snapshot)
        set_drive_link(source,staged_ppt,link)
        if staged_ppt.stat().st_size > LIMIT:
            raise ValueError('PPT exceeds 10 MB.')
        ppt_name, video_name = NAME+'.pptx', NAME+'.mp4'
        with zipfile.ZipFile(staged_zip,'w',zipfile.ZIP_DEFLATED) as z:
            z.write(staged_ppt,ppt_name)
            z.write(video,video_name)
        if staged_zip.stat().st_size > LIMIT:
            raise ValueError('Combined ZIP exceeds 10 MB. Reduce video size (aim below 9 MB) and retry.')
        with zipfile.ZipFile(staged_zip) as z:
            if z.testzip() is not None or set(z.namelist()) != {ppt_name,video_name}:
                raise ValueError('Submission ZIP integrity check failed.')
        staged_ppt.replace(SUB / ppt_name)
        target_video = SUB / video_name
        if video != target_video.resolve():
            shutil.copy2(video,target_video)
        staged_zip.replace(SUB / (NAME+'.zip'))
        status = {
            'status':'PPT and actual MP4 packaged; Drive URL inserted',
            'video_seconds':seconds,'video_bytes':target_video.stat().st_size,
            'ppt_bytes':(SUB/ppt_name).stat().st_size,'zip_bytes':(SUB/(NAME+'.zip')).stat().st_size,
            'drive_url':link,'drive_viewing_access':'User must test viewing access',
            'zip_sha256':hashlib.sha256((SUB/(NAME+'.zip')).read_bytes()).hexdigest(),
        }
        (SUB / 'Package_Status.json').write_text(json.dumps(status,indent=2),encoding='utf-8')
        pending = SUB / 'VIDEO_PENDING.txt'
        if pending.exists():
            pending.unlink()
        print('Prepared PPT, MP4 and ZIP in:',SUB)
        print(f'Video: {seconds:.2f} seconds; ZIP: {status["zip_bytes"]:,} bytes.')
        print('Test Drive viewing access, confirm names/IDs, watch the full video and check portal/deadline before the single submission.')
    finally:
        for p in staged:
            if p.exists():
                p.unlink()

if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print('NOT PACKAGED:',str(e),file=sys.stderr)
        sys.exit(1)
